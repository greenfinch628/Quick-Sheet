"""File adapters; Excel support is an optional, lazily loaded dependency."""
import csv
import json
import re
from pathlib import Path
from datetime import date, datetime
from sheet import Sheet, coordinates, address

def excel_module():
    try: import openpyxl
    except ImportError as error: raise ValueError('XLSX support requires openpyxl. Install with: python -m pip install -r requirements.txt') from error
    return openpyxl

def read_sheet(path):
    path=Path(path); ext=path.suffix.lower()
    if ext=='.mini': return Sheet.from_data(json.loads(path.read_text(encoding='utf-8-sig')))
    sheet=Sheet()
    if ext in ('.csv','.tsv'):
        with path.open(encoding='utf-8-sig',newline='') as stream:
            for r,row in enumerate(csv.reader(stream,delimiter='\t' if ext=='.tsv' else ',')):
                for c,value in enumerate(row): sheet.set(r,c,value)
        return sheet
    if ext!='.xlsx': raise ValueError('Use .mini, .csv, .tsv, or .xlsx')
    ox=excel_module(); wb=ox.load_workbook(path,data_only=False)
    try:
        ws=wb.active
        if ws.max_row*ws.max_column>1000000: raise ValueError('Excel sheet exceeds the lightweight import limit (1,000,000 cells)')
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None and not cell.has_style: continue
                key=cell.coordinate; value=cell.value
                if isinstance(value,(date,datetime)): value=value.isoformat(sep=' ') if isinstance(value,datetime) else value.isoformat()
                if value is not None: sheet.cells[key]=str(value)
                st={}
                fmt=cell.number_format
                if cell.is_date: st['format']='Date/time' if any(x in fmt.lower() for x in ('h','s')) else 'Date'
                elif '%' in fmt: st['format']='Percentage'
                elif '$' in fmt: st['format']='Currency'
                elif fmt=='@': st['format']='Text'
                elif fmt!='General': st['format']='Decimal'
                if '.' in fmt: st['decimals']=len(re.match(r'[0#]*',fmt.split('.',1)[1])[0])
                st.update(bold=bool(cell.font.bold),italic=bool(cell.font.italic))
                if cell.alignment.horizontal in ('left','center','right'): st['alignment']=cell.alignment.horizontal
                if cell.font.color and cell.font.color.type=='rgb': st['foreground']='#'+cell.font.color.rgb[-6:]
                if cell.fill.fill_type=='solid' and cell.fill.fgColor.type=='rgb': st['background']='#'+cell.fill.fgColor.rgb[-6:]
                sheet.styles[key]=st
        for key,dim in ws.column_dimensions.items():
            if dim.width: sheet.col_widths[str(coordinates(key+'1')[1])]=round(dim.width*7+5)
        for key,dim in ws.row_dimensions.items():
            if dim.height: sheet.row_heights[str(key-1)]=round(dim.height*4/3)
        sheet.table=ws.freeze_panes=='A2'
        return sheet
    finally: wb.close()

def write_sheet(sheet,path):
    path=Path(path); ext=path.suffix.lower()
    temp=path.with_name(path.stem+'.tmp'+ext)
    try:
        if ext=='.mini': temp.write_text(json.dumps(sheet.dump(),indent=2),encoding='utf-8')
        elif ext in ('.csv','.tsv'):
            with temp.open('w',encoding='utf-8-sig',newline='') as stream:
                writer=csv.writer(stream,delimiter='\t' if ext=='.tsv' else ',')
                rows,cols=sheet.extent()
                for r in range(rows): writer.writerow([sheet.display(r,c) for c in range(cols)])
        elif ext=='.xlsx':
            ox=excel_module(); wb=ox.Workbook(); ws=wb.active; ws.title='Quick-Sheet'
            from openpyxl.styles import Font, Alignment, PatternFill
            for key in sheet.cells.keys()|sheet.styles.keys():
                r,c=coordinates(key); raw=sheet.raw(r,c); st=sheet.style(r,c)
                target=ws.cell(r+1,c+1)
                if raw.startswith('=') and st.get('format')!='Text':
                    target.value=re.sub(r'^=\s*AUTOSUM(?=\s*\()', '=SUM',raw,flags=re.IGNORECASE)
                else:
                    value=sheet.value(r,c)
                    target.value=value if value!='' else None
                    if isinstance(value,str): target.data_type='s'
                decimals=int(st.get('decimals',0)); digits='.'+'0'*decimals if decimals else ''
                target.number_format={'Text':'@','Decimal':'#,##0'+digits,'Currency':'$#,##0'+digits,'Percentage':'0'+digits+'%','Date':'yyyy-mm-dd','Date/time':'yyyy-mm-dd hh:mm:ss'}.get(st.get('format'),'0'+digits)
                target.font=Font(name=sheet.cell_font,size=sheet.cell_size,bold=st.get('bold',False),italic=st.get('italic',False),color=st.get('foreground','#192c3b').lstrip('#'))
                target.alignment=Alignment(horizontal=st.get('alignment','left'))
                if st.get('background'): target.fill=PatternFill('solid',fgColor=st['background'].lstrip('#'))
            for key,value in sheet.col_widths.items(): ws.column_dimensions[address(0,int(key))[:-1]].width=max(1,(value-5)/7)
            for key,value in sheet.row_heights.items(): ws.row_dimensions[int(key)+1].height=value*3/4
            if sheet.table: ws.freeze_panes='A2'
            if sheet.table: ws.auto_filter.ref=ws.dimensions
            wb.save(temp); wb.close()
        else: raise ValueError('Use .mini, .csv, .tsv, or .xlsx')
        temp.replace(path)
    finally:
        if temp.exists(): temp.unlink()
